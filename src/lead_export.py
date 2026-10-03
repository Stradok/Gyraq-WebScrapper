"""Export scraped businesses as a cold-call list.

Within each search, the highest-rated businesses are the benchmark and the
low-rated ones are the targets. For every target we list its weaknesses
(compared with the benchmark, plus problems raised in its fresh bad Google
reviews and in Reddit / Quora / complaint-site mentions), the service that
fixes each one (website, AI voice agent, ...), and a call script. Everything
is rule-based on the scraped data - no LLM - and every comparison only uses
businesses from the same search.

Formats: PDF, Excel (.xlsx), CSV and JSON. The PDF is rendered by Chromium,
so any script (e.g. Urdu or Arabic business names) prints correctly.
"""
import csv
import html
import io
import json
import re
from datetime import datetime, timezone
from urllib.parse import parse_qsl, urlencode, urlparse, urlunparse

from . import config
from .company_profile import get_company_profile
from .results_store import read_result_file

KINDS = ("phones", "emails", "both")
FORMATS = ("pdf", "xlsx", "csv", "json")
TIERS = ("low", "all")
DEFAULT_MAX_RATING = 4.0

# Services we can offer, so the weakness -> fix mapping stays consistent.
WEBSITE = "Website"
WEBSITE_REDESIGN = "Website redesign"
VOICE = "AI voice agent"
CHAT = "WhatsApp / chat assistant"
BOOKING = "Online booking"
REPUTATION = "Reputation & reviews"
GBP = "Google Business Profile & local SEO"

# Businesses where the phone is the main way customers reach them.
CALL_DRIVEN = re.compile(
    r"dent|clinic|doctor|physio|hospital|salon|spa|barber|beauty|lawyer|solicitor|attorney|"
    r"plumb|electric|hvac|repair|garage|mechanic|locksmith|cleaner|cleaning|estate|vet|veterinar|"
    r"accountant|consult|gym|school|tutor|taxi|restaurant",
    re.I,
)

# What customers complain about in bad reviews / forum posts / complaint sites.
# (key, label, service-or-None, pattern, fix-or-None). A theme without a
# service is still shown (it tells you what the business is struggling with)
# but isn't something we sell a fix for.
THEMES = [
    (
        "website", "Website / design", WEBSITE_REDESIGN,
        re.compile(
            r"(web ?site|web ?page|their site|online (booking|form|portal|system|scheduling|ordering|payment))"
            r"[^.!?]{0,80}(outdated|old|confusing|slow|broken|doesn'?t work|does not work|not working|hard to|"
            r"difficult|terrible|awful|bad|horrible|unusable|crash|glitch|error|useless|impossible|clunky)|"
            r"(outdated|confusing|slow|broken|terrible|awful|poor|useless|clunky|unusable)[^.!?]{0,40}(web ?site|web ?page)|"
            r"can'?t (book|schedule|order|pay) online|no online (booking|ordering|scheduling)|"
            r"(web ?site|online form)[^.!?]{0,40}(didn'?t|doesn'?t|never) (work|load|submit)",
            re.I,
        ),
        "Redesign the website: clear layout, fast, mobile-friendly, with online booking and an easy contact form",
    ),
    (
        "calls", "Calls go unanswered", VOICE,
        re.compile(
            r"(didn'?t|did not|never|not|no one|nobody|couldn'?t|can'?t|cannot|unable to)\s+(\w+\s+)?"
            r"(answer|pick up|respond|reach|get through)|"
            r"(phone|call)s?\s+(is |are |was |were )?(not|never|unanswered|ignored)|"
            r"hard to reach|unreachable|voicemail|on hold|no answer|never picks?|engaged tone|busy line|"
            r"went straight to voicemail",
            re.I,
        ),
        "An AI voice agent answers every call 24/7, takes the details and books the visit, so no caller reaches voicemail",
    ),
    (
        "response", "No reply to messages and enquiries", CHAT,
        re.compile(
            r"slow(ly)? (to )?(reply|respond|response)|no (reply|response)|never (replied|responded|got back|called back)|"
            r"unresponsive|ignored (my |our )?(message|email|enquir|inquir)|poor communication|"
            r"didn'?t (get back|call back|return)|waiting (for )?(a |the )?(reply|response|callback|call back)|"
            r"lack of communication|no communication",
            re.I,
        ),
        "A WhatsApp / chat assistant replies to every enquiry within seconds, at any time of day",
    ),
    (
        "booking", "Booking, scheduling and no-shows", BOOKING,
        re.compile(
            r"(appointment|booking|reservation)s? (was |were )?(cancel|mix|problem|issue|hard|difficult|impossible|late|missed)|"
            r"couldn'?t (book|get an appointment)|no availability|long wait|waited (for )?(hours|ages|long|over)|"
            r"waiting (time|list)|overbook|double.?book|no.?show|didn'?t show( up)?|never showed|stood (me|us) up|"
            r"late (arrival|to the)|showed up late|arrived (very )?late|missed (the |my |our )?appointment|rescheduled",
            re.I,
        ),
        "Online booking with automatic reminders and arrival updates cuts no-shows, double bookings and waiting",
    ),
    (
        "staff", "Rude or unprofessional staff", None,
        re.compile(r"rude|unprofessional|disrespect|attitude|arrogant|dismissive|condescending|yelled|argument", re.I),
        None,
    ),
    (
        "price", "Overcharging / pricing", None,
        re.compile(r"over.?charg|rip.?off|too expensive|overpriced|hidden (fee|charge|cost)|price gouging|quoted (me )?(over|\$)|scam", re.I),
        None,
    ),
    (
        "quality", "Poor workmanship / quality", None,
        re.compile(r"poor (work|quality|job)|shoddy|botched|had to (redo|call someone else)|still (leaking|broken|not working)|didn'?t fix|never fixed|made it worse|wasn'?t fixed", re.I),
        None,
    ),
]


# How the problem sounds when said out loud on a call.
SPOKEN = {
    "website": "customers are complaining that your website is hard to use",
    "calls": "customers say your calls go unanswered",
    "response": "customers say they can't get a reply from you",
    "booking": "customers are complaining about booking and waiting",
}


# ---------- small helpers ----------

def _digits(phone: str | None) -> str:
    return re.sub(r"\D", "", phone or "")


def _clean_url(url: str | None) -> str | None:
    """Drop tracking parameters (utm_*) so links read cleanly."""
    if not url:
        return url
    try:
        p = urlparse(url)
        q = [(k, v) for k, v in parse_qsl(p.query, keep_blank_values=True) if not k.lower().startswith("utm_")]
        return urlunparse(p._replace(query=urlencode(q)))
    except Exception:
        return url


def _stars(b: dict) -> str:
    r = b.get("rating")
    return f"{r}★" if r is not None else "unrated"


def _one_line(b: dict) -> str:
    site = "has a website" if b.get("website") else "no website"
    return f"{b.get('name')} ({_stars(b)}, {b.get('review_count') or 0} reviews, {site})"


def _quote(text: str, n: int = 140) -> str:
    return "“" + re.sub(r"\s+", " ", text).strip()[:n] + "”"


# ---------- benchmark & evidence ----------

def benchmarks(peers: list[dict], limit: int = 3) -> list[dict]:
    """Best businesses in a search: highest rating among those with enough
    reviews to be credible (20+), falling back to most-reviewed."""
    rated = [p for p in peers if p.get("rating") is not None and (p.get("review_count") or 0) >= 20]
    if rated:
        rated.sort(key=lambda p: (p["rating"], p.get("review_count") or 0), reverse=True)
        return rated[:limit]
    others = sorted((p for p in peers if p.get("review_count")), key=lambda p: p["review_count"], reverse=True)
    return others[:limit]


def evidence_items(biz: dict) -> list[dict]:
    """Everything negative said about the business: fresh bad Google reviews,
    older low-rated reviews, and negative Reddit / Quora / complaint-site hits."""
    items, seen = [], set()
    for r in list(biz.get("negative_reviews") or []) + list(biz.get("reviews") or []):
        text = (r.get("text") or "").strip()
        if not text or (r.get("rating") is not None and r["rating"] > 3):
            continue
        key = text[:60]
        if key in seen:
            continue
        seen.add(key)
        items.append({
            "source": "Google review", "text": text, "rating": r.get("rating"),
            "age_days": r.get("age_days"), "when": r.get("relative_time") or "", "url": "",
        })
    for m in biz.get("mentions") or []:
        if m.get("negative"):
            items.append({
                "source": m.get("source") or "Web", "text": f"{m.get('title', '')}. {m.get('snippet', '')}".strip(". "),
                "rating": None, "age_days": None, "when": "", "url": m.get("url") or "",
            })
    items.sort(key=lambda i: i["age_days"] if i["age_days"] is not None else 10**6)
    return items


def classify_themes(items: list[dict]) -> list[dict]:
    """[{key, label, service, fix, count, items}] for themes that appear, most common first."""
    out = []
    for key, label, service, pattern, fix in THEMES:
        hits = [i for i in items if pattern.search(i["text"])]
        if hits:
            out.append({"key": key, "label": label, "service": service, "fix": fix, "count": len(hits), "items": hits})
    out.sort(key=lambda t: t["count"], reverse=True)
    return out


def _evidence_line(i: dict) -> str:
    if i["source"] == "Google review":
        parts = [p for p in (f"{int(i['rating'])}★" if i.get("rating") else "", i["when"]) if p]
        return f"Google review ({', '.join(parts)}): {_quote(i['text'])}" if parts else f"Google review: {_quote(i['text'])}"
    link = f" ({i['url']})" if i["url"] else ""
    return f"{i['source']}: {_quote(i['text'], 120)}{link}"


# ---------- analysis ----------

def analyze(biz: dict, peers: list[dict], max_rating: float = DEFAULT_MAX_RATING) -> dict:
    """{gaps: [...], themes: [...], items: [...]}. Each gap: weakness, evidence, service, fix, kind."""
    gaps: list[dict] = []
    top = benchmarks([p for p in peers if p is not biz])
    best = top[0] if top else None
    avg_reviews = sum(t.get("review_count") or 0 for t in top) / len(top) if top else None
    reviews = biz.get("review_count") or 0
    rating = biz.get("rating")

    def add(weakness, evidence, service, fix, kind="data"):
        gaps.append({"weakness": weakness, "evidence": evidence, "service": service, "fix": fix, "kind": kind})

    if not biz.get("website"):
        ev = f"{best['name']}, the top-rated in this search, has a website." if best and best.get("website") else ""
        add(
            "No website. People who find them on Google Maps have nowhere to learn more, book or enquire",
            ev, WEBSITE,
            "Build a fast mobile-friendly website with click-to-call and an enquiry/booking form",
        )
    elif not biz.get("email"):
        add(
            "Their website shows no email address (phone or contact form only), so every enquiry depends on someone picking up",
            "", CHAT,
            "Add instant follow-up on every enquiry through a chat assistant, WhatsApp and email",
        )

    if rating is not None and (rating < max_rating or (best and best.get("rating") and best["rating"] - rating >= 0.6)):
        ev = f"Top-rated in this search: {best['name']} at {best['rating']}★." if best and best.get("rating") else ""
        add(
            f"Rated {rating}★, which pushes customers to higher-rated competitors",
            ev, REPUTATION,
            "Ask every happy customer for a review automatically and reply to negative reviews professionally",
        )
    elif rating is None:
        add(
            "No rating yet, so customers have no proof the business is good",
            "", REPUTATION,
            "Start a review-collection system (QR code plus a WhatsApp/SMS request after each job)",
        )

    if avg_reviews and rating is not None and reviews < max(20, 0.5 * avg_reviews):
        add(
            f"Only {reviews} reviews against about {int(avg_reviews)} for the top-rated in this search",
            _one_line(best) if best else "", REPUTATION,
            "Run a review-collection system (QR code plus a WhatsApp/SMS request after each job)",
        )
    elif not avg_reviews and rating is not None and reviews < 20:
        add(f"Only {reviews} reviews, which is little social proof", "", REPUTATION,
            "Run a review-collection system (QR code plus a WhatsApp/SMS request after each job)")

    if not biz.get("hours"):
        add(
            "No opening hours on their Google listing, so customers can't tell if they're open",
            "", GBP, "Complete and optimise their Google Business Profile (hours, photos, services, posts)",
        )

    items = evidence_items(biz)
    themes = classify_themes(items)
    for t in themes:
        if not t["service"]:
            continue
        freshest = t["items"][0]
        fresh = [i for i in t["items"] if i["age_days"] is not None and i["age_days"] <= config.FRESH_REVIEW_DAYS]
        how = f"{t['count']} bad review/mention(s)" + (f", {len(fresh)} within the last year" if fresh else "")
        add(
            f"{t['label']}: customers complain about this ({how})",
            _evidence_line(freshest), t["service"], t["fix"], kind="review",
        )
        gaps[-1]["spoken"] = SPOKEN.get(t["key"], t["label"].lower())

    if gaps and not any(g["service"] == VOICE for g in gaps) and CALL_DRIVEN.search(biz.get("category") or ""):
        add(
            "Customers reach this kind of business by phone, and a missed call is a lost job",
            "Based on the type of business, not on a review", VOICE,
            "An AI voice agent answers every call 24/7, takes details and books the visit",
            kind="opportunity",
        )
    # Lead with the most persuasive point: no website, then what customers
    # themselves complain about, then rating/reviews, and the weak ones last.
    def weight(g):
        if g["service"] == WEBSITE:
            return 0
        if g["kind"] == "review":
            return 1
        if g["kind"] == "opportunity":
            return 4
        if g["service"] == CHAT:
            return 3
        return 2

    gaps.sort(key=weight)
    return {"gaps": gaps, "themes": themes, "items": items}


def priority(biz: dict, gaps: list[dict]) -> tuple[float, str]:
    """Bigger is a better prospect: more weaknesses and a lower rating."""
    rating = biz.get("rating")
    score = len(gaps) + (max(0.0, 4.5 - rating) * 2 if rating is not None else 3)
    score = round(min(score, 10), 1)
    return score, ("High" if score >= 6 else "Medium" if score >= 3 else "Low")


def _better_rivals(biz: dict, peers: list[dict]) -> list[dict]:
    return [p for p in benchmarks([p for p in peers if p is not biz]) if (p.get("rating") or 0) >= (biz.get("rating") or 0)]


def comparison(biz: dict, peers: list[dict]) -> str:
    rivals = _better_rivals(biz, peers)
    if not rivals:
        return ""
    return f"Best in this search: {_one_line(rivals[0])} vs this lead: {_one_line(biz)}"


def call_script(biz: dict, gaps: list[dict], peers: list[dict], company: dict) -> str:
    us = company.get("company_name") or "our team"
    name = biz.get("name") or "your business"
    category = (biz.get("category") or "business").lower()
    if not gaps:
        return (
            f"Hi, is this {name}? I'm calling from {us}. You're doing well online compared to other "
            f"{category}s nearby, so I'll keep it short. We help strong businesses turn that visibility "
            "into more booked jobs. Could I send you a two-minute overview? "
            "What's the best email or WhatsApp number for you?"
        )
    rivals = _better_rivals(biz, peers)
    rival = rivals[0] if rivals else None
    first = gaps[0]
    headline = first.get("spoken") or first["weakness"].split(". ")[0].rstrip(".")
    lines = [
        f"Hi, is this {name}? I'm calling from {us}. I'll be brief. I was looking at {category}s in your "
        "area on Google and noticed something that's costing you customers.",
        f"Right now: {headline[:1].lower() + headline[1:]}.",
    ]
    if rival:
        lines.append(
            f"Meanwhile {rival['name']} is showing up with {rival.get('review_count') or 0} reviews at "
            f"{rival.get('rating')}★{' and a proper website' if rival.get('website') else ''}. "
            "When someone searches, they get the call, and you lose that job."
        )
    review_gap = next((g for g in gaps if g["kind"] == "review"), None)
    if review_gap:
        lines.append(f"Customers are saying it themselves: {review_gap['evidence']}")
    services = list(dict.fromkeys(g["service"] for g in gaps))[:3]
    lines.append(f"We can fix this: {', '.join(services)}. {first['fix']}.")
    lines.append(
        "Could I send you a short breakdown of exactly where you stand against your competitors? "
        "What's the best email or WhatsApp number for you?"
    )
    return "\n".join(lines)


# ---------- collecting ----------

def collect(files: list[str], kind: str, tier: str = "all", max_rating: float = DEFAULT_MAX_RATING):
    """Returns (rows, benchmark_rows, meta). Rows are best prospects first."""
    company = get_company_profile()
    rows, bench_rows, seen, queries = [], [], set(), []
    for file in files:
        data = read_result_file(file)
        if not data:
            continue
        peers = data.get("results") or []
        query = data.get("query")
        queries.append(query)
        for b in benchmarks(peers):
            bench_rows.append({
                "search_query": query, "name": b.get("name"), "category": b.get("category"),
                "rating": b.get("rating"), "review_count": b.get("review_count"),
                "website": _clean_url(b.get("website")), "phone": b.get("phone"), "address": b.get("address"),
            })
        for biz in peers:
            phone, email = biz.get("phone"), (biz.get("email") or "").strip()
            if kind == "phones" and not phone:
                continue
            if kind == "emails" and not email:
                continue
            if kind == "both" and not (phone or email):
                continue
            if tier == "low" and biz.get("rating") is not None and biz["rating"] >= max_rating:
                continue
            key = _digits(phone) or email.lower()
            if key in seen:
                continue
            seen.add(key)
            a = analyze(biz, peers, max_rating)
            gaps, themes, items = a["gaps"], a["themes"], a["items"]
            score, label = priority(biz, gaps)
            fresh_bad = [i for i in items if i["source"] == "Google review"][:3]
            rows.append({
                "priority": label,
                "score": score,
                "name": biz.get("name"),
                "category": biz.get("category"),
                "phone": phone if kind != "emails" else None,
                "email": (email or None) if kind != "phones" else None,
                "website": _clean_url(biz.get("website")),
                "address": biz.get("address"),
                "rating": biz.get("rating"),
                "review_count": biz.get("review_count"),
                "search_query": query,
                "vs_top_rated": comparison(biz, peers),
                "weaknesses": " | ".join(g["weakness"] for g in gaps),
                "review_themes": ", ".join(f"{t['label']} ({t['count']})" for t in themes),
                "fresh_bad_reviews": " | ".join(_evidence_line(i) for i in fresh_bad),
                "online_research": " | ".join(
                    f"{m.get('source')}{' (negative)' if m.get('negative') else ''}: {m.get('url')}"
                    for m in (biz.get("mentions") or [])
                ),
                "evidence": " | ".join(g["evidence"] for g in gaps if g["evidence"]),
                "services_to_offer": ", ".join(dict.fromkeys(g["service"] for g in gaps)),
                "how_we_help": " | ".join(g["fix"] for g in gaps),
                "call_script": call_script(biz, gaps, peers, company),
                "gaps": gaps,
                "themes": [{"label": t["label"], "count": t["count"]} for t in themes],
                "mentions": biz.get("mentions") or [],
                "fresh_reviews": fresh_bad,
            })
    rows.sort(key=lambda r: r["score"], reverse=True)
    demand: dict[str, int] = {}
    for r in rows:
        for s in dict.fromkeys(g["service"] for g in r["gaps"]):
            demand[s] = demand.get(s, 0) + 1
    meta = {
        "company": company.get("company_name") or "Gyraq",
        "queries": queries, "kind": kind, "tier": tier, "max_rating": max_rating,
        "generated": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
        "demand": dict(sorted(demand.items(), key=lambda kv: kv[1], reverse=True)),
    }
    return rows, bench_rows, meta


COLUMNS = [
    ("priority", "Priority"), ("score", "Score"), ("name", "Business"), ("category", "Category"),
    ("phone", "Phone"), ("email", "Email"), ("website", "Website"), ("address", "Address"),
    ("rating", "Rating"), ("review_count", "Reviews"), ("search_query", "Search"),
    ("vs_top_rated", "Vs top-rated"), ("weaknesses", "Weaknesses"), ("review_themes", "Review themes"),
    ("fresh_bad_reviews", "Fresh bad reviews"), ("online_research", "Online research (Reddit, Quora, complaint sites)"),
    ("services_to_offer", "Services to offer"), ("how_we_help", "How we can help"),
    ("call_script", "Call script"),
]


# ---------- writers ----------

def _csv(rows) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([label for _, label in COLUMNS])
    for r in rows:
        w.writerow([r.get(k) if r.get(k) is not None else "" for k, _ in COLUMNS])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel reads UTF-8


def _xlsx(rows, bench_rows, meta) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="C8102E")
    head_font = Font(bold=True, color="FFFFFF")
    wrap = Alignment(wrap_text=True, vertical="top")

    def style_header(ws, widths):
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for c in ws[1]:
            c.fill, c.font, c.alignment = head_fill, head_font, Alignment(vertical="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    ws = wb.active
    ws.title = "Leads"
    ws.append([label for _, label in COLUMNS])
    widths = [10, 7, 28, 18, 18, 26, 30, 30, 8, 9, 22, 44, 56, 34, 60, 50, 30, 56, 70]
    colors = {"High": "F8D7DA", "Medium": "FFF3CD", "Low": "E2E3E5"}
    for r in rows:
        ws.append([r.get(k) if r.get(k) is not None else "" for k, _ in COLUMNS])
        ws.cell(row=ws.max_row, column=1).fill = PatternFill("solid", fgColor=colors[r["priority"]])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = wrap
    style_header(ws, widths)

    bs = wb.create_sheet("Top-rated benchmarks")
    bs.append(["Search", "Business", "Category", "Rating", "Reviews", "Website", "Phone", "Address"])
    for b in bench_rows:
        bs.append([b["search_query"], b["name"], b["category"], b["rating"], b["review_count"],
                   b["website"] or "none", b["phone"], b["address"]])
    style_header(bs, [26, 30, 20, 8, 9, 30, 18, 36])

    sm = wb.create_sheet("Summary", 0)
    sm.append(["Gyraq lead report", ""])
    sm["A1"].font = Font(bold=True, size=14, color="C8102E")
    sm.append(["Generated", meta["generated"]])
    sm.append(["Searches", ", ".join(q for q in meta["queries"] if q)])
    sm.append(["Businesses shown", f"low-rated only (below {meta['max_rating']}★)" if meta["tier"] == "low" else "all"])
    sm.append(["Leads in this report", len(rows)])
    sm.append([])
    sm.append(["Service to offer", "Leads who need it"])
    for c in sm[7]:
        c.font = Font(bold=True)
    for service, n in meta["demand"].items():
        sm.append([service, n])
    sm.column_dimensions["A"].width = 34
    sm.column_dimensions["B"].width = 70

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _pdf(rows, bench_rows, meta) -> bytes:
    """Plain list, one block per client with bold labels, links clickable."""
    e = html.escape

    def line(label: str, value, link: bool = False) -> str:
        if value in (None, ""):
            return ""
        text = e(str(value))
        if link:
            text = f'<a href="{e(str(value), quote=True)}">{text}</a>'
        return f"<div><b>{e(label)}:</b> {text}</div>"

    def multi(label: str, items: list[str]) -> str:
        items = [i for i in items if i]
        if not items:
            return ""
        return f"<div><b>{e(label)}:</b></div>" + "".join(f"<div class='ind'>{i}</div>" for i in items)

    blocks = []
    for n, r in enumerate(rows, 1):
        issues = [e(g["weakness"]) + (f"<br><i>{e(g['evidence'])}</i>" if g["evidence"] else "") for g in r["gaps"]]
        helps = [f"<b>{e(g['service'])}:</b> {e(g['fix'])}" for g in r["gaps"]]
        fresh = [e(_evidence_line(i)) for i in r["fresh_reviews"]]
        research = []
        for m in r["mentions"]:
            url = m.get("url") or ""
            tag = " (negative)" if m.get("negative") else ""
            research.append(f"{e(str(m.get('source')))}{tag}: <a href=\"{e(url, quote=True)}\">{e(url)}</a>")
        themes = ", ".join(f"{t['label']} ({t['count']})" for t in r["themes"])
        rating = r["rating"] if r["rating"] is not None else "No rating"
        blocks.append(
            f"<div class='client'><h3>CLIENT {n}</h3>"
            + line("Client Name", r["name"])
            + line("Category", r["category"])
            + line("Client Number", r["phone"] or "Not provided")
            + line("Email", r["email"] or "Not provided")
            + (line("Website", r["website"], link=True) if r["website"] else line("Website", "None found"))
            + line("Address", r["address"])
            + line("Rating", rating)
            + line("Review Count", r["review_count"])
            + line("Search Query", r["search_query"])
            + line("Priority", f"{r['priority']} ({r['score']}/10)")
            + multi("Issues", issues or ["No clear weaknesses found."])
            + line("Review Themes", themes)
            + multi("Fresh Bad Reviews", fresh)
            + multi("Online Research", research)
            + line("Services We Can Offer", r["services_to_offer"])
            + multi("How We Can Help", helps)
            + line("Competitor Benchmark", r["vs_top_rated"])
            + f"<div><b>Call Script:</b> {e(r['call_script']).replace(chr(10), '<br>')}</div>"
            + "<hr></div>"
        )

    demand = "".join(f"<div class='ind'>{e(s)}: {n}</div>" for s, n in meta["demand"].items())
    scope = f"low-rated businesses only (below {meta['max_rating']}★)" if meta["tier"] == "low" else "all businesses"
    doc = f"""<!doctype html><html><head><meta charset="utf-8"><style>
  @page {{ size: A4; margin: 16mm 14mm; }}
  body {{ font-family: "Times New Roman", "Liberation Serif", "Noto Serif", "DejaVu Serif", serif; color: #000; font-size: 11.5px; line-height: 1.35; }}
  h1 {{ text-align: center; font-size: 22px; margin: 0 0 14px; }}
  h3 {{ font-size: 14px; margin: 14px 0 6px; }}
  .ind {{ margin-left: 14px; }}
  a {{ color: #1a56c4; }}
  hr {{ border: 0; border-top: 1px solid #000; margin: 14px 0 0; }}
</style></head><body>
<h1>Contactable Contact List</h1>
<div>Total Clients: {len(rows)}</div>
<div>All client information is arranged individually. Website links are clickable.</div>
<div>Searches: {e(', '.join(q for q in meta['queries'] if q))} · {scope} · {e(meta['generated'])}</div>
{('<div><b>Services these clients need:</b></div>' + demand) if demand else ''}
{''.join(blocks)}
</body></html>"""
    from playwright.sync_api import sync_playwright

    with sync_playwright() as p:
        browser = p.chromium.launch(args=["--no-sandbox"])
        try:
            page = browser.new_page()
            page.set_content(doc)
            return page.pdf(format="A4", print_background=True)
        finally:
            browser.close()


def build_export(
    files: list[str], kind: str, fmt: str, tier: str = "all", max_rating: float = DEFAULT_MAX_RATING
) -> tuple[bytes, str, str]:
    rows, bench_rows, meta = collect(files, kind, tier, max_rating)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    name = f"leads_{kind}_{stamp}.{fmt}"
    if fmt == "json":
        return json.dumps(rows, ensure_ascii=False, indent=2).encode(), name, "application/json"
    if fmt == "xlsx":
        return (
            _xlsx(rows, bench_rows, meta), name,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    if fmt == "pdf":
        return _pdf(rows, bench_rows, meta), name, "application/pdf"
    return _csv(rows), name, "text/csv"
